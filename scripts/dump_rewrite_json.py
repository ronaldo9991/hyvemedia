"""Dump the rewrite workbook to a UTF-8 JSON file so we can read clean Unicode
strings (em-dashes, curly quotes) without console encoding interference.
"""
import json
from openpyxl import load_workbook

wb = load_workbook("content-audit-HYVE-rewrite-v2.xlsx", data_only=True)
ws = wb["Content Audit"]

rows = []
for r in range(2, ws.max_row + 1):
    rows.append({
        "id": (ws.cell(row=r, column=1).value or "").strip(),
        "section": (ws.cell(row=r, column=2).value or "").strip(),
        "location": (ws.cell(row=r, column=3).value or "").strip(),
        "source": (ws.cell(row=r, column=4).value or "").strip(),
        "current": ws.cell(row=r, column=5).value or "",
        "new": ws.cell(row=r, column=6).value or "",
        "notes": (ws.cell(row=r, column=7).value or "").strip(),
    })

with open("scripts/rewrite.json", "w", encoding="utf-8") as f:
    json.dump(rows, f, ensure_ascii=False, indent=2)

print(f"Wrote scripts/rewrite.json with {len(rows)} rows.")
