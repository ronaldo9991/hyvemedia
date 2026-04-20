"""Dump rewrite-v3 workbook to UTF-8 JSON + also write a diff-vs-v2 summary
so we can tell what the content team changed this round.
"""
import json
from openpyxl import load_workbook

wb3 = load_workbook("content-audit-HYVE-rewrite-v3.xlsx", data_only=True)
ws3 = wb3["Content Audit"]

rows = []
for r in range(2, ws3.max_row + 1):
    rows.append({
        "id": (ws3.cell(row=r, column=1).value or "").strip(),
        "section": (ws3.cell(row=r, column=2).value or "").strip(),
        "location": (ws3.cell(row=r, column=3).value or "").strip(),
        "source": (ws3.cell(row=r, column=4).value or "").strip(),
        "current": ws3.cell(row=r, column=5).value or "",
        "new": ws3.cell(row=r, column=6).value or "",
        "notes": (ws3.cell(row=r, column=7).value or "").strip(),
    })

with open("scripts/rewrite_v3.json", "w", encoding="utf-8") as f:
    json.dump(rows, f, ensure_ascii=False, indent=2)

# Diff against v2 so we can focus edits on what actually changed this round.
try:
    with open("scripts/rewrite.json", encoding="utf-8") as f:
        v2 = {row["id"]: row for row in json.load(f)}
except FileNotFoundError:
    v2 = {}

diffs = []
for row in rows:
    v2_new = (v2.get(row["id"], {}) or {}).get("new", "")
    if (row["new"] or "") != (v2_new or ""):
        diffs.append({
            "id": row["id"],
            "section": row["section"],
            "location": row["location"],
            "source": row["source"],
            "v2_new": v2_new,
            "v3_new": row["new"],
            "v3_notes": row["notes"],
        })

with open("scripts/rewrite_v3_diff.json", "w", encoding="utf-8") as f:
    json.dump(diffs, f, ensure_ascii=False, indent=2)

print(f"v3 rows: {len(rows)}  |  changes vs v2: {len(diffs)}")
