"""Generate content-audit.xlsx — a review sheet of every piece of user-facing
copy on the HYVE Media marketing site, with a blank column for rewrites.

Each row has a stable ID. Keep the ID when handing the filled-in sheet back
so edits can be applied deterministically.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ---------- Content inventory ----------
# (section, location, file:line, current)
ROWS = [
    # --- <head> (index.html) ---
    ("Meta / SEO", "Browser tab title", "index.html:7", "HYVE Media | An Onchain Communications Firm"),
    ("Meta / SEO", "Meta description", "index.html:9", "HYVE Media is an onchain communications firm helping Web3 teams win with strategy, thought leadership, press relations, and content."),
    ("Meta / SEO", "Meta author", "index.html:12", "HYVE Media"),
    ("Meta / SEO", "OG site_name", "index.html:19", "HYVE Media"),
    ("Meta / SEO", "OG title", "index.html:20", "HYVE Media | An Onchain Communications Firm"),
    ("Meta / SEO", "OG description", "index.html:23", "We help Web3 teams spotlight their message with strategy, thought leadership, press relations, and content."),
    ("Meta / SEO", "OG image alt", "index.html:27", "HYVE Media branding"),
    ("Meta / SEO", "Twitter title", "index.html:30", "HYVE Media | An Onchain Communications Firm"),
    ("Meta / SEO", "Twitter description", "index.html:33", "HYVE Media helps onchain teams earn visibility and trust through high-impact communications."),

    # --- Navbar ---
    ("Navbar", "Left link 1", "Navbar.tsx:7", "Why"),
    ("Navbar", "Left link 2", "Navbar.tsx:8", "What"),
    ("Navbar", "Left link 3", "Navbar.tsx:9", "How"),
    ("Navbar", "Left link 4", "Navbar.tsx:10", "Promise"),
    ("Navbar", "Right link 1", "Navbar.tsx:14", "Join"),
    ("Navbar", "Contact link", "Navbar.tsx:77", "Contact"),
    ("Navbar", "Contact email (mailto)", "Navbar.tsx:73", "hello@hyvemedia.com"),
    ("Navbar", "Logo alt text (scrolled)", "Navbar.tsx:55", "HYVE Logo"),
    ("Navbar", "Logo alt text (top)", "Navbar.tsx:57", "HYVE Icon"),

    # --- Hero ---
    ("Hero", "Tagline", "HeroSection.tsx:38", "An Onchain Communications Firm"),
    ("Hero", "Main headline (wordmark)", "HeroSection.tsx:53", "HYVE Media"),

    # --- Storytelling ---
    ("Storytelling", "Section headline", "StorytellingSection.tsx:51", "Storytelling is a long game. We're here to help you win it."),

    # --- Why We Exist ---
    ("Why We Exist", "Eyebrow", "WhySection.tsx:20", "Why We Exist"),
    ("Why We Exist", "Headline", "WhySection.tsx:31", "We believe in the long-term impact of decentralized technology and the power of communications."),
    ("Why We Exist", "Body paragraph", "WhySection.tsx:42", "Since 2016, we have focused exclusively on this innovative industry, helping teams on the cutting edge tell their stories. Whether it's infrastructure, the convergence of blockchain & AI, or venture capital — we do what we do because without great storytelling, no one would know this technology exists."),
    ("Why We Exist", "Image alt", "WhySection.tsx:55", "Why we exist visual"),

    # --- Logo Marquee ---
    ("Logo Marquee", "Eyebrow", "LogoMarquee.tsx:20", "Regularly published in"),
    ("Logo Marquee", "Publication 1", "LogoMarquee.tsx:2", "Bloomberg"),
    ("Logo Marquee", "Publication 2", "LogoMarquee.tsx:2", "CoinDesk"),
    ("Logo Marquee", "Publication 3", "LogoMarquee.tsx:2", "CNBC"),
    ("Logo Marquee", "Publication 4", "LogoMarquee.tsx:2", "Forbes"),
    ("Logo Marquee", "Publication 5", "LogoMarquee.tsx:2", "The Block"),
    ("Logo Marquee", "Publication 6", "LogoMarquee.tsx:2", "Fortune"),
    ("Logo Marquee", "Publication 7", "LogoMarquee.tsx:3", "Bitcoin Magazine"),
    ("Logo Marquee", "Publication 8", "LogoMarquee.tsx:3", "Decrypt"),
    ("Logo Marquee", "Publication 9", "LogoMarquee.tsx:3", "MarketWatch"),
    ("Logo Marquee", "Publication 10", "LogoMarquee.tsx:3", "BeInCrypto"),

    # --- Published Photo ---
    ("Published Photo", "Image alt", "PublishedPhotoSection.tsx:16", "Media publication feature visual"),

    # --- What We Do ---
    ("What We Do", "Eyebrow", "WhatWeDoSection.tsx:67", "What We Do"),
    ("What We Do", "Headline line 1", "WhatWeDoSection.tsx:70", "You bring the technology."),
    ("What We Do", "Headline line 2", "WhatWeDoSection.tsx:72", "We spotlight your message on center stage."),
    ("What We Do", "Service 1 — title", "WhatWeDoSection.tsx:5", "Strategy"),
    ("What We Do", "Service 1 — description", "WhatWeDoSection.tsx:5", "Refine your narrative to support your business goals."),
    ("What We Do", "Service 2 — title", "WhatWeDoSection.tsx:7", "Thought Leadership"),
    ("What We Do", "Service 2 — description", "WhatWeDoSection.tsx:8", "Build the profiles of your team's leaders to broaden the impact of your story in the media."),
    ("What We Do", "Service 3 — title", "WhatWeDoSection.tsx:11", "Press Relations"),
    ("What We Do", "Service 3 — description", "WhatWeDoSection.tsx:12", "Work our relationship magic to garner the coverage your project deserves."),
    ("What We Do", "Service 4 — title", "WhatWeDoSection.tsx:15", "Content"),
    ("What We Do", "Service 4 — description", "WhatWeDoSection.tsx:16", "Craft compelling content that meets your target audience where they're at."),

    # --- How We Work ---
    ("How We Work", "Eyebrow", "HowWeWorkSection.tsx:205", "How We Work"),
    ("How We Work", "Headline part 1", "HowWeWorkSection.tsx:215", "Five precise moves."),
    ("How We Work", "Headline part 2 (accent)", "HowWeWorkSection.tsx:216", "Zero filler."),
    ("How We Work", "Step 1 — phase", "HowWeWorkSection.tsx:27", "01 · Discover"),
    ("How We Work", "Step 1 — title", "HowWeWorkSection.tsx:28", "Discovery & Immersion"),
    ("How We Work", "Step 1 — blurb", "HowWeWorkSection.tsx:29", "We learn your brand inside-out and find the story worth telling."),
    ("How We Work", "Step 1 — meta tags", "HowWeWorkSection.tsx:30", "Audit · Audience · Signals"),
    ("How We Work", "Step 2 — phase", "HowWeWorkSection.tsx:34", "02 · Strategy"),
    ("How We Work", "Step 2 — title", "HowWeWorkSection.tsx:35", "Strategy & Blueprint"),
    ("How We Work", "Step 2 — blurb", "HowWeWorkSection.tsx:36", "A sharp comms plan with clear themes, channels and KPIs."),
    ("How We Work", "Step 2 — meta tags", "HowWeWorkSection.tsx:37", "Messaging · Channels · KPIs"),
    ("How We Work", "Step 3 — phase", "HowWeWorkSection.tsx:41", "03 · Create"),
    ("How We Work", "Step 3 — title", "HowWeWorkSection.tsx:42", "Creative Development"),
    ("How We Work", "Step 3 — blurb", "HowWeWorkSection.tsx:43", "High-conviction content that makes your message unmissable."),
    ("How We Work", "Step 3 — meta tags", "HowWeWorkSection.tsx:44", "Content · Thought Leadership"),
    ("How We Work", "Step 4 — phase", "HowWeWorkSection.tsx:48", "04 · Launch"),
    ("How We Work", "Step 4 — title", "HowWeWorkSection.tsx:49", "Activation & Launch"),
    ("How We Work", "Step 4 — blurb", "HowWeWorkSection.tsx:50", "Owned and earned rollout, orchestrated for real momentum."),
    ("How We Work", "Step 4 — meta tags", "HowWeWorkSection.tsx:51", "Rollout · Media · Monitoring"),
    ("How We Work", "Step 5 — phase", "HowWeWorkSection.tsx:55", "05 · Scale"),
    ("How We Work", "Step 5 — title", "HowWeWorkSection.tsx:56", "Optimize & Scale"),
    ("How We Work", "Step 5 — blurb", "HowWeWorkSection.tsx:57", "We double down on what works, then compound it month over month."),
    ("How We Work", "Step 5 — meta tags", "HowWeWorkSection.tsx:58", "Insights · Iteration · Scale"),
    ("How We Work", "Closing tag line", "HowWeWorkSection.tsx:253", "Precision. Not performance theatre."),

    # --- Promise ---
    ("Our Promise", "Eyebrow", "PromiseSection.tsx:38", "Our Promise To You"),
    ("Our Promise", "Headline line 1", "PromiseSection.tsx:44", "You can lean on us as your communications experts."),
    ("Our Promise", "Headline line 2", "PromiseSection.tsx:46", "\"We've got you!\""),
    ("Our Promise", "Promise 1 — title", "PromiseSection.tsx:7", "Extremely communicative"),
    ("Our Promise", "Promise 1 — description", "PromiseSection.tsx:8", "You'll hear from us A LOT. We won't leave you waiting or wondering what's next. Communication is the name of the game for us."),
    ("Our Promise", "Promise 2 — title", "PromiseSection.tsx:12", "High quality results"),
    ("Our Promise", "Promise 2 — description", "PromiseSection.tsx:13", "The proof is in the pudding — we deliver results that create the greatest impact."),
    ("Our Promise", "Promise 3 — title", "PromiseSection.tsx:17", "Crypto obsessed"),
    ("Our Promise", "Promise 3 — description", "PromiseSection.tsx:18", "Communicating the value of cutting-edge companies and projects to the world gives us purpose; it's a mission that's personal to us."),
    ("Our Promise", "Promise 4 — title", "PromiseSection.tsx:22", "Win the long game"),
    ("Our Promise", "Promise 4 — description", "PromiseSection.tsx:23", "Your goals are our goals. We're here to help you go the distance and realize your project's vision."),

    # --- Join ---
    ("Join / Careers", "Eyebrow", "JoinSection.tsx:86", "Join Our Team"),
    ("Join / Careers", "Headline", "JoinSection.tsx:92", "Manifesting your dream job at HYVE Media?"),
    ("Join / Careers", "Intro paragraph (full, with link text 'email')", "JoinSection.tsx:98", "Well, of course you are. If any of the following sounds up your alley, jump on the next flight to LA. Or actually, maybe just start with an email."),
    ("Join / Careers", "Intro link mailto", "JoinSection.tsx:101", "hello@hyvemedia.com"),
    ("Join / Careers", "Value 1 — title", "JoinSection.tsx:60", "Radiate positivity"),
    ("Join / Careers", "Value 1 — description (two paragraphs)", "JoinSection.tsx:61",
     "At HYVE Media, it's not what we do that makes us memorable — it's how we do it. Core to that philosophy is how we show up in the world: radiating positivity in everything we do.\n\nWhether it's our \"can do\" attitude or the optimistic energy we bring to each interaction, you can feel our warm & welcoming spirit across every touchpoint. It's no wonder clients are naturally drawn to our magnetic presence."),
    ("Join / Careers", "Value 2 — title", "JoinSection.tsx:65", "Crazy empathetic"),
    ("Join / Careers", "Value 2 — description", "JoinSection.tsx:66", "Not to get all woo woo, but people always come first. We listen, are accountable, and put ourselves in our colleagues' and clients' shoes daily. But not in any kind of manufactured way — authenticity, always."),
    ("Join / Careers", "Value 3 — title", "JoinSection.tsx:70", "High quality, always"),
    ("Join / Careers", "Value 3 — description (two paragraphs)", "JoinSection.tsx:71",
     "Clients love us not only because we're warm, but because we produce at a world-class rate. Among our partners we're known as the gold standard of media. And tbh, that's probably underselling it.\n\nThat level of quality is represented across our brand and the people we hire. We are a dependable, high-caliber team that knows how to win the long game."),

    # --- Promise Photos ---
    ("Promise Photos", "Image 1 alt", "PromisePhotosSection.tsx:18", "Promise visual one"),
    ("Promise Photos", "Image 2 alt", "PromisePhotosSection.tsx:35", "Promise visual two"),

    # --- Contact ---
    ("Contact", "Eyebrow", "ContactSection.tsx:21", "Contact Us"),
    ("Contact", "Headline", "ContactSection.tsx:24", "Let's build your communications edge."),
    ("Contact", "Body paragraph", "ContactSection.tsx:27", "Share your goals and we'll reach out with a tailored plan. This is a demo form."),
    ("Contact", "Form label — Full Name", "ContactSection.tsx:48", "Full Name"),
    ("Contact", "Form placeholder — Full Name", "ContactSection.tsx:52", "Jane Doe"),
    ("Contact", "Form label — Email", "ContactSection.tsx:64", "Email"),
    ("Contact", "Form placeholder — Email", "ContactSection.tsx:68", "name@company.com"),
    ("Contact", "Form label — Company", "ContactSection.tsx:80", "Company"),
    ("Contact", "Form placeholder — Company", "ContactSection.tsx:84", "HYVE Labs"),
    ("Contact", "Form label — Service Focus", "ContactSection.tsx:96", "Service Focus"),
    ("Contact", "Form select placeholder", "ContactSection.tsx:108", "Select a service"),
    ("Contact", "Form select — option 1", "ContactSection.tsx:110", "Strategy"),
    ("Contact", "Form select — option 2", "ContactSection.tsx:111", "Thought Leadership"),
    ("Contact", "Form select — option 3", "ContactSection.tsx:112", "Press Relations"),
    ("Contact", "Form select — option 4", "ContactSection.tsx:113", "Content"),
    ("Contact", "Form label — Message", "ContactSection.tsx:119", "Message"),
    ("Contact", "Form placeholder — Message", "ContactSection.tsx:123", "Tell us about your goals..."),
    ("Contact", "Submit button", "ContactSection.tsx:143", "Send Message"),

    # --- Footer ---
    ("Footer", "Logo alt", "Footer.tsx:11", "HYVE Media"),
    ("Footer", "Tagline", "Footer.tsx:19", "Made in Dubai"),
    ("Footer", "Address label", "Footer.tsx:25", "Address"),
    ("Footer", "Address (two lines)", "Footer.tsx:27", "Office 2904, Marina Plaza,\nDubai Marina, Dubai, UAE"),
    ("Footer", "Navigate label", "Footer.tsx:34", "Navigate"),
    ("Footer", "Nav link 1", "Footer.tsx:36", "Why"),
    ("Footer", "Nav link 2", "Footer.tsx:36", "What"),
    ("Footer", "Nav link 3", "Footer.tsx:36", "How"),
    ("Footer", "Nav link 4", "Footer.tsx:36", "Promise"),
    ("Footer", "Nav link 5", "Footer.tsx:36", "Join"),
    ("Footer", "Connect label", "Footer.tsx:50", "Connect"),
    ("Footer", "Instagram URL", "Footer.tsx:52", "https://instagram.com/hyvemedia"),
    ("Footer", "Instagram link text", "Footer.tsx:53", "Instagram"),
    ("Footer", "LinkedIn URL", "Footer.tsx:55", "https://linkedin.com/company/hyvemedia"),
    ("Footer", "LinkedIn link text", "Footer.tsx:56", "LinkedIn"),
    ("Footer", "Contact email", "Footer.tsx:58", "hello@hyvemedia.com"),
    ("Footer", "Copyright line", "Footer.tsx:70", "© 2026 HYVE Media. All rights reserved."),

    # --- WhatsApp Float ---
    ("WhatsApp Float", "Prefilled message", "WhatsAppFloat.tsx:4", "Hi HYVE Media, I would like to learn more about your communications services."),
    ("WhatsApp Float", "WhatsApp number (in URL)", "WhatsAppFloat.tsx:7", "971500000000"),
    ("WhatsApp Float", "aria-label", "WhatsAppFloat.tsx:14", "Chat on WhatsApp"),

    # --- 404 ---
    ("404 Page", "Headline", "NotFound.tsx:15", "404"),
    ("404 Page", "Body", "NotFound.tsx:16", "Oops! Page not found"),
    ("404 Page", "Return link text", "NotFound.tsx:18", "Return to Home"),

    # --- Unused components (not currently on the live page) ---
    ("[Unused] Team", "Eyebrow", "TeamSection.tsx:15", "Meet The Dream Team"),
    ("[Unused] Team", "Headline", "TeamSection.tsx:17", "Collaboration Maxis"),
    ("[Unused] Team", "Member 1 — name", "TeamSection.tsx:4", "Kelley W."),
    ("[Unused] Team", "Member 1 — role", "TeamSection.tsx:4", "CEO"),
    ("[Unused] Team", "Member 1 — bio", "TeamSection.tsx:4", "A pioneer in this industry for over 9 years. Believes in synchronicity and the power of positivity."),
    ("[Unused] Team", "Member 2 — name", "TeamSection.tsx:5", "Alex G."),
    ("[Unused] Team", "Member 2 — role", "TeamSection.tsx:5", "Vice President"),
    ("[Unused] Team", "Member 2 — bio", "TeamSection.tsx:5", "11 years in emerging tech PR, with a passion for creative problem-solving and strategy."),
    ("[Unused] Team", "Member 3 — name", "TeamSection.tsx:6", "Eunice H."),
    ("[Unused] Team", "Member 3 — role", "TeamSection.tsx:6", "Vice President"),
    ("[Unused] Team", "Member 3 — bio", "TeamSection.tsx:6", "A lifelong communications professional whose love for tech brought her to Web3."),
    ("[Unused] Team", "Member 4 — name", "TeamSection.tsx:7", "Jean N."),
    ("[Unused] Team", "Member 4 — role", "TeamSection.tsx:7", "Senior Account Executive"),
    ("[Unused] Team", "Member 4 — bio", "TeamSection.tsx:7", "Loves serving as a daily point of contact. At her core, she's a theater kid bringing creative energy."),
    ("[Unused] Team", "Member 5 — name", "TeamSection.tsx:8", "Dean P."),
    ("[Unused] Team", "Member 5 — role", "TeamSection.tsx:8", "Account Executive"),
    ("[Unused] Team", "Member 5 — bio", "TeamSection.tsx:8", "Combines ten years in web3 with a lifelong passion for writing and communication."),
    ("[Unused] Team", "Member 6 — name", "TeamSection.tsx:9", "Kayla K."),
    ("[Unused] Team", "Member 6 — role", "TeamSection.tsx:9", "Operations Coordinator"),
    ("[Unused] Team", "Member 6 — bio", "TeamSection.tsx:9", "Passionate about connecting with people and building deep, meaningful relationships."),

    ("[Unused] Podcast", "Eyebrow", "PodcastSection.tsx:15", "Podcast"),
    ("[Unused] Podcast", "Show title", "PodcastSection.tsx:17", "Proof-Of-PR"),
    ("[Unused] Podcast", "Intro paragraph", "PodcastSection.tsx:20", "Your passport to the world of onchain communications. Get an inside look into the future of finance, technology, and entrepreneurship."),
    ("[Unused] Podcast", "Episode 10 title", "PodcastSection.tsx:4", "Bitcoin and The Art of Social Strategies"),
    ("[Unused] Podcast", "Episode 9 title", "PodcastSection.tsx:5", "The Intersection of Culture and Blockchain"),
    ("[Unused] Podcast", "Episode 8 title", "PodcastSection.tsx:6", "The Symbiotic Relationship Between PR and Journalism"),
    ("[Unused] Podcast", "Episode 7 title", "PodcastSection.tsx:7", "Covering Bitcoin and Crypto Adoption in Real Life"),
    ("[Unused] Podcast", "Episode 6 title", "PodcastSection.tsx:8", "Marketing in the Web3 Space and Transitioning from Web2"),
    ("[Unused] Podcast", "Episode 5 title", "PodcastSection.tsx:9", "Using PR and Twitter to Drive a Strong Media Presence"),
]


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Content Audit"

    headers = ["ID", "Section", "Location / Element", "Source (file:line)", "Current Content", "New Content", "Notes"]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="262626")
    thin = Side(style="thin", color="D4D4D4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="left")
        cell.border = border

    new_content_fill = PatternFill("solid", fgColor="FFF3DC")
    unused_fill = PatternFill("solid", fgColor="F1F1F1")
    wrap = Alignment(wrap_text=True, vertical="top")

    for i, (section, location, source, current) in enumerate(ROWS, start=1):
        row_num = i + 1
        ws.cell(row=row_num, column=1, value=f"R{i:03d}")
        ws.cell(row=row_num, column=2, value=section)
        ws.cell(row=row_num, column=3, value=location)
        ws.cell(row=row_num, column=4, value=source)
        ws.cell(row=row_num, column=5, value=current)
        ws.cell(row=row_num, column=6, value="")  # to be filled by content team
        ws.cell(row=row_num, column=7, value="")  # notes column

        is_unused = section.startswith("[Unused]")
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=row_num, column=col)
            c.alignment = wrap
            c.border = border
            if col == 6:
                c.fill = new_content_fill
            elif is_unused:
                c.fill = unused_fill

    widths = [8, 22, 36, 32, 70, 70, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # --- Instructions sheet ---
    ins = wb.create_sheet("How to use")
    ins.column_dimensions["A"].width = 110
    lines = [
        ("HYVE Media — Content Audit", True),
        ("", False),
        ("Purpose", True),
        ("Every user-facing string on the marketing site has been extracted to the 'Content Audit' tab. Your team rewrites anything that should change in the highlighted 'New Content' column. Leave the cell blank to keep the current copy as-is.", False),
        ("", False),
        ("Workflow", True),
        ("1. Send this workbook to the content team.", False),
        ("2. They fill in the yellow 'New Content' column for any row they want to change. Leave it blank to keep the current copy.", False),
        ("3. They can use the 'Notes' column to flag questions, tone decisions, or translation context.", False),
        ("4. Send the workbook back — the 'ID' column is how edits are mapped onto the codebase deterministically, so please keep it intact.", False),
        ("", False),
        ("Conventions", True),
        ("• A blank 'New Content' cell = no change.", False),
        ("• Line breaks inside a single string are preserved — use Alt+Enter (Windows) / Option+Enter (Mac) to add a newline in Excel.", False),
        ("• Straight quotes and em-dashes (—) are fine; they render as-is.", False),
        ("• Rows tagged '[Unused]' (grey) belong to components that exist in the code but are NOT currently rendered on the live page (TeamSection, PodcastSection). Only edit these if you plan to bring those sections back.", False),
        ("", False),
        ("Section order on the live page", True),
        ("Navbar → Hero → Storytelling → Why We Exist → Logo Marquee → Published Photo → What We Do → How We Work → Our Promise → Join → Promise Photos → Contact → Footer (+ floating WhatsApp button)", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ins.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            c.font = Font(bold=True, size=12)

    out = "content-audit.xlsx"
    wb.save(out)
    print(f"Wrote {out} — {len(ROWS)} content rows across {len({r[0] for r in ROWS})} sections.")


if __name__ == "__main__":
    build()
