"""Dump every row of the rewrite workbook where 'New Content' is non-empty,
so we can see exactly what edits to apply.
"""
import sys
from openpyxl import load_workbook

wb = load_workbook("content-audit-HYVE-rewrite-v2.xlsx", data_only=True)
ws = wb["Content Audit"]

def norm(v):
    if v is None:
        return ""
    return str(v)

print(f"Total rows (excl header): {ws.max_row - 1}")
print()

changes = []
notes_only = []

for r in range(2, ws.max_row + 1):
    rid = norm(ws.cell(row=r, column=1).value).strip()
    section = norm(ws.cell(row=r, column=2).value).strip()
    location = norm(ws.cell(row=r, column=3).value).strip()
    source = norm(ws.cell(row=r, column=4).value).strip()
    current = norm(ws.cell(row=r, column=5).value)
    new = norm(ws.cell(row=r, column=6).value)
    notes = norm(ws.cell(row=r, column=7).value).strip()

    has_new = new.strip() != ""
    has_notes = notes != ""

    if has_new:
        changes.append((rid, section, location, source, current, new, notes))
    elif has_notes:
        notes_only.append((rid, section, location, source, current, notes))

print(f"Rows with new content: {len(changes)}")
print(f"Rows with notes only:  {len(notes_only)}")
print()
print("=" * 100)
print("CHANGES")
print("=" * 100)
for rid, section, location, source, current, new, notes in changes:
    print(f"\n[{rid}] {section} — {location}  ({source})")
    print(f"  OLD: {current!r}")
    print(f"  NEW: {new!r}")
    if notes:
        print(f"  NOTE: {notes}")

if notes_only:
    print()
    print("=" * 100)
    print("NOTES ONLY (no rewrite)")
    print("=" * 100)
    for rid, section, location, source, current, notes in notes_only:
        print(f"\n[{rid}] {section} — {location}  ({source})")
        print(f"  CURRENT: {current!r}")
        print(f"  NOTE: {notes}")
